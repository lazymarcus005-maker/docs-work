"""Document parsers (ticket #5, spec §10, FR-003).

The registry produces canonical documents behind a replaceable interface
(NFR-004). Built-in parsers use lightweight CPU libraries; if Docling is
installed it is preferred automatically, and additional parsers can be
registered for new formats (spec §8).
"""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Callable

from .canonical import canonical_document

PARSER_VERSION = "builtin-1"
DOCLING_PARSER_VERSION = "docling-1"


def docling_available() -> bool:
    """True when the docling package is installed (cheap check, no import)."""
    import importlib.util

    try:
        return importlib.util.find_spec("docling") is not None
    except (ImportError, ValueError):
        return False


def docling_version() -> str | None:
    try:
        from importlib.metadata import version

        return version("docling")
    except Exception:  # noqa: BLE001 — package metadata simply absent
        return None


# parsers: (suffix, kind) -> callable(bytes, filename) -> list[element dicts]
_REGISTRY: dict[tuple[str, str], Callable] = {}


def register(suffixes: tuple[str, ...], kind: str):
    def deco(fn):
        for s in suffixes:
            _REGISTRY[(s, kind)] = fn
        return fn

    return deco


def supported_suffixes() -> set[str]:
    return {s for (s, _kind) in _REGISTRY}


def parse_file(document_id: str, path: Path, kind: str,
               use_docling: bool = True) -> dict:
    """Parse into a canonical document; the result is tagged with the parser
    that actually ran so incremental processing can track it (§11)."""
    suffix = path.suffix.lower()
    if use_docling and docling_available() and kind in ("pdf", "docx", "pptx"):
        docling_elements = _try_docling(path)
        if docling_elements is not None:
            doc = canonical_document(
                document_id, path.name, suffix.lstrip("."), docling_elements
            )
            doc["parser"] = DOCLING_PARSER_VERSION
            return doc
    fn = _REGISTRY.get((suffix, kind))
    if fn is None:
        raise ValueError(f"no parser registered for {suffix} ({kind})")
    elements = fn(path.read_bytes(), path.name)
    doc = canonical_document(document_id, path.name, suffix.lstrip("."), elements)
    doc["parser"] = PARSER_VERSION
    return doc


def _try_docling(path: Path) -> list[dict] | None:
    """Prefer Docling when installed (spec §10.1); None → fall back."""
    try:
        from docling.document_converter import DocumentConverter  # type: ignore
    except ImportError:
        return None
    try:
        result = DocumentConverter().convert(str(path))
        doc = result.document
        elements = []
        for item, _level in doc.iterate_items():
            text = getattr(item, "text", "") or ""
            label = str(getattr(item, "label", "") or "").lower()
            etype = "heading" if "heading" in label or "title" in label else (
                "table" if "table" in label else "paragraph"
            )
            if text or etype == "table":
                elements.append({"type": etype, "text": text})
        return elements
    except Exception:
        return None


# --------------------------------------------------------------- text / md
@register((".txt",), "text")
def parse_txt(data: bytes, filename: str) -> list[dict]:
    text = data.decode("utf-8", errors="replace")
    return [
        {"type": "paragraph", "text": para.strip()}
        for para in text.split("\n\n")
        if para.strip()
    ]


@register((".md", ".markdown"), "markdown")
def parse_markdown(data: bytes, filename: str) -> list[dict]:
    elements: list[dict] = []
    in_code = False
    table_buf: list[str] = []
    for line in data.decode("utf-8", errors="replace").splitlines():
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        is_table_line = stripped.startswith("|") and stripped.endswith("|")
        if not is_table_line and table_buf:
            elements.append({"type": "table", "text": "\n".join(table_buf)})
            table_buf = []
        if not stripped:
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            elements.append(
                {"type": "heading", "level": min(level, 6), "text": stripped.lstrip("#").strip()}
            )
        elif is_table_line:
            table_buf.append(stripped)
        elif stripped.startswith(("-", "*", "+")) or (
            stripped[:2] in ("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.")
        ):
            elements.append({"type": "list_item", "text": stripped.lstrip("-*+ ").strip()})
        else:
            elements.append({"type": "paragraph", "text": stripped})
    if table_buf:
        elements.append({"type": "table", "text": "\n".join(table_buf)})
    return elements


# --------------------------------------------------------------------- pdf
@register((".pdf",), "pdf")
def parse_pdf(data: bytes, filename: str) -> list[dict]:
    try:
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(data))
        elements: list[dict] = []
        total_chars = 0
        for pageno, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            total_chars += len(text)
            for para in [p.strip() for p in text.split("\n\n") if p.strip()]:
                elements.append({"type": "paragraph", "text": para, "page": pageno})
        pages = len(reader.pages) or 1
        density = total_chars / pages
        if pages >= 1 and density < 50:
            elements.append({
                "type": "image",
                "text": "",
                "meta": {"needs_ocr": True, "text_density": round(density, 1)},
            })
        return elements
    except Exception as e:
        raise ValueError(f"PDF parse failed: {e}") from e


# -------------------------------------------------------------------- docx
@register((".docx",), "docx")
def parse_docx(data: bytes, filename: str) -> list[dict]:
    try:
        import docx as docx_lib
    except ImportError as e:
        raise ValueError("python-docx not installed") from e

    df = docx_lib.Document(io.BytesIO(data))
    elements: list[dict] = []

    def walk(parent):
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in parent.iter_inner_content():
            if isinstance(child, Paragraph):
                style = (child.style.name or "").lower()
                text = child.text.strip()
                if not text:
                    continue
                if style.startswith("heading"):
                    try:
                        level = int(style.split()[-1])
                    except ValueError:
                        level = 1
                    elements.append({"type": "heading", "level": level, "text": text})
                elif style.startswith("list"):
                    elements.append({"type": "list_item", "text": text})
                else:
                    elements.append({"type": "paragraph", "text": text})
            elif isinstance(child, Table):
                elements.append({"type": "table", "text": _table_text(child)})

    walk(df)
    return elements


def _table_text(table) -> str:
    rows = []
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        rows.append(" | ".join(cells))
    return "\n".join(rows)


# -------------------------------------------------------------------- xlsx
@register((".xlsx",), "xlsx")
def parse_xlsx(data: bytes, filename: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    elements: list[dict] = []
    for sheet in wb.worksheets:
        elements.append({"type": "heading", "level": 1, "text": sheet.title})
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
            if len(rows) >= 500:
                break
        if rows:
            elements.append({"type": "table", "text": "\n".join(rows), "meta": {"sheet": sheet.title}})
    return elements


# -------------------------------------------------------------------- pptx
@register((".pptx",), "pptx")
def parse_pptx(data: bytes, filename: str) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    elements: list[dict] = []
    for i, slide in enumerate(prs.slides, start=1):
        elements.append({"type": "slide", "text": f"Slide {i}", "page": i})
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        elements.append({"type": "paragraph", "text": text, "page": i})
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    rows.append(" | ".join(c.text.strip() for c in row.cells))
                if rows:
                    elements.append({"type": "table", "text": "\n".join(rows), "page": i})
    return elements


# --------------------------------------------------------------------- csv
@register((".csv",), "csv")
def parse_csv(data: bytes, filename: str) -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = []
    for i, row in enumerate(reader):
        if i >= 1000:
            break
        if any(c.strip() for c in row):
            rows.append(" | ".join(c.strip() for c in row))
    return [{"type": "table", "text": "\n".join(rows)}] if rows else []


# ------------------------------------------------------------------ images
@register((".png", ".jpg", ".jpeg"), "image")
def parse_image(data: bytes, filename: str) -> list[dict]:
    # Text requires OCR (ticket #15); mark so the pipeline reports it properly.
    return [{"type": "image", "text": "", "meta": {"needs_ocr": True, "text_density": 0.0}}]

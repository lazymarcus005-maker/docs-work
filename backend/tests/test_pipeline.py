"""Ticket #5: parse → chunk → index pipeline + text search + job queue."""
from __future__ import annotations

import io
import time

import pytest

from app.documents.chunker import chunk_document
from app.storage import filesystem as fs


def _upload(client, pid, name, content, mime="application/octet-stream"):
    return client.post(
        f"/api/projects/{pid}/files",
        files={"files": (name, io.BytesIO(content), mime)},
    )


def _wait_status(client, pid, file_id, status, timeout=10.0):
    deadline = time.time() + timeout
    while time.time() < deadline:
        files = client.get(f"/api/projects/{pid}/files").json()["files"]
        doc = next(f for f in files if f["id"] == file_id)
        if doc["status"] == status:
            return doc
        time.sleep(0.1)
    raise AssertionError(f"file never reached {status}: {doc}")


def _make_docx_bytes(paragraphs):
    import docx as docx_lib

    from io import BytesIO

    d = docx_lib.Document()
    d.add_heading("Authentication", level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_pipeline_reaches_ready_and_searchable(client):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    content = _make_docx_bytes([
        "The gateway forwards authentication requests to the identity service.",
        "Tokens expire after 30 minutes and are stored in Redis.",
    ])
    res = _upload(client, pid, "SRS.docx", content,
                  "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    fid = res.json()["files"][0]["file_id"]

    doc = _wait_status(client, pid, fid, "READY")
    assert doc["status"] == "READY"

    search = client.post(
        f"/api/projects/{pid}/search", json={"query": "authentication gateway"}
    ).json()
    assert search["results"], "uploaded content must be searchable"
    top = search["results"][0]
    assert top["document_name"] == "SRS.docx"
    assert top["section_path"] == ["Authentication"]
    assert top["chunk_id"] and top["text"]

    # canonical document persisted to the project workspace
    parsed = fs.parsed_dir(client.app.state.settings.workspace_root, pid) / f"{fid}.json"
    assert parsed.exists()
    canonical = parsed.read_text()
    assert "Authentication" in canonical


def test_txt_pipeline_and_status_lifecycle(client):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    fid = _upload(client, pid, "notes.txt", b"Redis caches token validation results.").json()["files"][0]["file_id"]
    doc = _wait_status(client, pid, fid, "READY")
    assert doc["error"] is None
    search = client.post(f"/api/projects/{pid}/search", json={"query": "Redis"}).json()
    assert search["results"]


def test_search_finds_hyphenated_identifier(client):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    content = b"Synthetic test marker: COBALT-ORCHID-731."
    fid = _upload(client, pid, "identifiers.txt", content).json()["files"][0]["file_id"]
    _wait_status(client, pid, fid, "READY")

    search = client.post(
        f"/api/projects/{pid}/search", json={"query": "COBALT-ORCHID-731"}
    ).json()

    assert search["results"], "hyphenated identifiers should match indexed terms"
    assert "COBALT-ORCHID-731" in search["results"][0]["text"]


def test_scanned_pdf_fails_actionably_with_enable_ocr(client):
    import pypdf

    from io import BytesIO

    writer = pypdf.PdfWriter()
    for _ in range(2):
        writer.add_blank_page(width=612, height=792)  # no extractable text
    buf = BytesIO()
    writer.write(buf)

    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    fid = _upload(client, pid, "scan.pdf", buf.getvalue(), "application/pdf").json()["files"][0]["file_id"]
    doc = _wait_status(client, pid, fid, "FAILED")
    assert "OCR is disabled" in doc["error"]
    assert "Enable OCR" in doc["error"]


def test_failed_file_does_not_block_others(client):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    _upload(client, pid, "broken.pdf", b"%PDF-1.4 garbage", "application/pdf")
    good = _upload(client, pid, "good.txt", b"searchable gateway content").json()["files"][0]["file_id"]
    _wait_status(client, pid, good, "READY")
    files = client.get(f"/api/projects/{pid}/files").json()["files"]
    statuses = {f["status"] for f in files}
    assert "READY" in statuses  # project remains usable


def test_processing_endpoint_reports_lifecycle(client):
    pid = client.post("/api/projects", json={"name": "P"}).json()["id"]
    fid = _upload(client, pid, "a.txt", b"some gateway auth content here").json()["files"][0]["file_id"]
    _wait_status(client, pid, fid, "READY")
    status = client.get(f"/api/projects/{pid}/processing").json()
    assert status["files"][0]["status"] == "READY"
    assert all(j["status"] == "SUCCEEDED" for j in status["latest_jobs"])


# ------------------------------------------------------------------- queue
def test_queue_priority_order(settings):
    from app import db
    from app.jobs import queue

    conn = db.connect(settings.db_path)
    db.init_db(conn)
    pid = "prj_x"
    conn.execute(
        "INSERT INTO projects (id, name, created_at, updated_at) VALUES (?, 'P', 't', 't')",
        (pid,),
    )
    queue.enqueue(conn, pid, "EXTRACT_ENTITIES", priority=queue.PRIORITY_ENRICHMENT)
    queue.enqueue(conn, pid, "PARSE_DOCUMENT", priority=queue.PRIORITY_PARSE)
    job = queue.claim_next(conn, "w")
    assert job["type"] == "PARSE_DOCUMENT" and job["status"] == "RUNNING"

    # interrupted RUNNING jobs requeue on restart (spec §44)
    assert queue.requeue_interrupted(conn) == 1
    again = queue.claim_next(conn, "w")
    assert again["id"] == job["id"]
    queue.finish_job(conn, again["id"])
    assert queue.claim_next(conn, "w")["type"] == "EXTRACT_ENTITIES"
    conn.close()


# ------------------------------------------------------------------ chunker
def test_chunker_xlsx_sheet_paths():
    import openpyxl

    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "APIs"
    ws.append(["endpoint", "method"])
    ws.append(["/auth/login", "POST"])
    buf = BytesIO()
    wb.save(buf)

    from app.documents.parser import parse_file

    doc = parse_file("doc_x", None, "xlsx") if False else None
    # parse via bytes path
    from app.documents.parser import parse_xlsx

    elements = parse_xlsx(buf.getvalue(), "api.xlsx")
    chunks = chunk_document({"document_id": "doc_x", "source": "api.xlsx", "elements": elements})
    assert any(c["section_path"] == ["APIs"] and "/auth/login" in c["text"] for c in chunks)
